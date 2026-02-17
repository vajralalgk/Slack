#!/usr/bin/env python3
"""
Generate ECTP Architecture & Future Roadmap Word Document (.docx)

This script produces a professionally designed Word document with embedded
matplotlib diagrams for the Enterprise Cloud Transformation Platform roadmap.

Output: /home/user/Slack/docs/word-documents/ECTP-Architecture-Future-Roadmap.docx
"""

import os
import io
import math
import tempfile
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import matplotlib.patheffects as pe
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
import numpy as np

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
DEEP_BLUE = "#1B3A5C"
TEAL = "#2596BE"
ACCENT_GOLD = "#E8A838"
SUCCESS_GREEN = "#27AE60"
ALERT_RED = "#E74C3C"
PURPLE = "#8E44AD"
LIGHT_BLUE = "#D6EAF8"
WHITE = "#FFFFFF"
LIGHT_GRAY = "#F2F3F4"
DARK_GRAY = "#2C3E50"

DEEP_BLUE_RGB = RGBColor(0x1B, 0x3A, 0x5C)
TEAL_RGB = RGBColor(0x25, 0x96, 0xBE)
GOLD_RGB = RGBColor(0xE8, 0xA8, 0x38)
GREEN_RGB = RGBColor(0x27, 0xAE, 0x60)
RED_RGB = RGBColor(0xE7, 0x4C, 0x3C)
PURPLE_RGB = RGBColor(0x8E, 0x44, 0xAD)
WHITE_RGB = RGBColor(0xFF, 0xFF, 0xFF)
BLACK_RGB = RGBColor(0x00, 0x00, 0x00)
DARK_GRAY_RGB = RGBColor(0x2C, 0x3E, 0x50)

OUTPUT_DIR = "/home/user/Slack/docs/word-documents"
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "ECTP-Architecture-Future-Roadmap.docx")
TEMP_DIR = tempfile.mkdtemp(prefix="ectp_diagrams_")

DPI = 300

# ---------------------------------------------------------------------------
# Diagram helper: save figure to a temporary PNG
# ---------------------------------------------------------------------------

def save_figure(fig, name):
    path = os.path.join(TEMP_DIR, f"{name}.png")
    fig.savefig(path, dpi=DPI, bbox_inches="tight", facecolor=fig.get_facecolor(), edgecolor="none")
    plt.close(fig)
    return path


# ---------------------------------------------------------------------------
# Diagram 1: Roadmap Timeline (Gantt-style swimlane)
# ---------------------------------------------------------------------------

def create_roadmap_timeline():
    fig, ax = plt.subplots(figsize=(14, 6))
    fig.set_facecolor(WHITE)
    ax.set_facecolor("#F8FAFC")

    phases = [
        ("Phase 1: Foundation", "Jan 2026", "Apr 2026", 0, 4, DEEP_BLUE, "$115K"),
        ("Phase 2: Integration", "May 2026", "Sep 2026", 4, 5, TEAL, "$150K"),
        ("Phase 3: Optimization", "Oct 2026", "Jan 2027", 9, 4, SUCCESS_GREEN, "$80K"),
        ("Phase 4: Automation", "Feb 2027", "Jun 2027", 13, 5, PURPLE, "$100K"),
        ("Phase 5: Innovation", "Jul 2027", "Dec 2027", 18, 6, ACCENT_GOLD, "$85K"),
    ]

    milestones = [
        (4, "Foundation\nComplete", DEEP_BLUE),
        (9, "Integration\nComplete", TEAL),
        (13, "Optimization\nComplete", SUCCESS_GREEN),
        (18, "Automation\nComplete", PURPLE),
        (24, "Roadmap\nComplete", ALERT_RED),
    ]

    for i, (name, start, end, offset, duration, color, budget) in enumerate(phases):
        y = len(phases) - i - 1
        bar = ax.barh(y, duration, left=offset, height=0.6, color=color, alpha=0.85,
                       edgecolor="white", linewidth=1.5, zorder=3)
        # Label inside bar
        ax.text(offset + duration / 2, y, f"{name}\n{budget}",
                ha="center", va="center", fontsize=8, fontweight="bold",
                color="white", zorder=4)

    # Milestones
    for month, label, color in milestones:
        ax.plot(month, -0.8, marker="D", markersize=10, color=color, zorder=5)
        ax.text(month, -1.25, label, ha="center", va="top", fontsize=6,
                color=color, fontweight="bold")

    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep",
              "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun",
              "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    ax.set_xticks(range(24))
    ax.set_xticklabels(months, fontsize=7, rotation=45)
    ax.set_yticks(range(len(phases)))
    ax.set_yticklabels([p[0] for p in reversed(phases)], fontsize=9, fontweight="bold")

    # Year labels
    ax.text(5.5, 5.6, "2026", ha="center", fontsize=12, fontweight="bold", color=DEEP_BLUE)
    ax.text(17.5, 5.6, "2027", ha="center", fontsize=12, fontweight="bold", color=DEEP_BLUE)
    ax.axvline(x=12, color=DEEP_BLUE, linestyle="--", alpha=0.4, linewidth=1)

    ax.set_xlim(-0.5, 24.5)
    ax.set_ylim(-2, 5.8)
    ax.set_title("ECTP Roadmap Timeline  |  Jan 2026 - Dec 2027  |  Total Budget: $530K",
                 fontsize=13, fontweight="bold", color=DEEP_BLUE, pad=12)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(axis="x", alpha=0.15)
    fig.tight_layout()
    return save_figure(fig, "01_roadmap_timeline")


# ---------------------------------------------------------------------------
# Diagram 2: Phase Budget Distribution (Donut)
# ---------------------------------------------------------------------------

def create_budget_donut():
    fig, ax = plt.subplots(figsize=(9, 6))
    fig.set_facecolor(WHITE)

    labels = ["Phase 1\nFoundation", "Phase 2\nIntegration", "Phase 3\nOptimization",
              "Phase 4\nAutomation", "Phase 5\nInnovation"]
    sizes = [115, 150, 80, 100, 85]
    colors = [DEEP_BLUE, TEAL, SUCCESS_GREEN, PURPLE, ACCENT_GOLD]
    explode = (0.03, 0.03, 0.03, 0.03, 0.03)
    total = sum(sizes)

    wedges, texts, autotexts = ax.pie(
        sizes, labels=None, autopct="", startangle=140, colors=colors,
        explode=explode, pctdistance=0.78, wedgeprops=dict(width=0.45, edgecolor="white", linewidth=2))

    for i, (w, s) in enumerate(zip(wedges, sizes)):
        angle = (w.theta2 + w.theta1) / 2
        x = 0.72 * math.cos(math.radians(angle))
        y = 0.72 * math.sin(math.radians(angle))
        pct = s / total * 100
        ax.text(x, y, f"${s}K\n{pct:.0f}%", ha="center", va="center",
                fontsize=9, fontweight="bold", color="white")

    # Legend
    legend_labels = [f"{l.replace(chr(10), ' ')} - ${s}K ({s/total*100:.1f}%)"
                     for l, s in zip(labels, sizes)]
    ax.legend(wedges, legend_labels, loc="center left", bbox_to_anchor=(1.0, 0.5),
              fontsize=9, frameon=True, fancybox=True, shadow=True)

    ax.text(0, 0, f"Total\n${total}K", ha="center", va="center",
            fontsize=16, fontweight="bold", color=DEEP_BLUE)
    ax.set_title("Phase Budget Distribution", fontsize=14, fontweight="bold", color=DEEP_BLUE, pad=16)
    fig.tight_layout()
    return save_figure(fig, "02_budget_donut")


# ---------------------------------------------------------------------------
# Diagram 3: Phase 1 Deliverables Progress
# ---------------------------------------------------------------------------

def create_deliverables_progress():
    fig, ax = plt.subplots(figsize=(12, 6))
    fig.set_facecolor(WHITE)
    ax.set_facecolor("#F8FAFC")

    deliverables = [
        ("1.1  Architecture Design & ADRs", 100, "Complete"),
        ("1.2  Core FastAPI Platform", 100, "Complete"),
        ("1.3  Health Check Endpoints", 100, "Complete"),
        ("1.4  Terraform Modules", 60, "In Progress"),
        ("1.5  CI/CD Pipeline", 10, "Planned"),
        ("1.6  Docker Build & ECR", 10, "Planned"),
        ("1.7  ServiceNow Integration", 5, "Planned"),
        ("1.8  Ellucian Ethos API", 5, "Planned"),
        ("1.9  Security Baseline", 0, "Planned"),
        ("1.10 Production Deployment", 0, "Planned"),
    ]

    status_colors = {"Complete": SUCCESS_GREEN, "In Progress": ACCENT_GOLD, "Planned": LIGHT_BLUE}
    status_text_colors = {"Complete": "white", "In Progress": "white", "Planned": DARK_GRAY}

    y_pos = np.arange(len(deliverables))
    names = [d[0] for d in reversed(deliverables)]
    pcts = [d[1] for d in reversed(deliverables)]
    statuses = [d[2] for d in reversed(deliverables)]

    for i, (name, pct, status) in enumerate(zip(names, pcts, statuses)):
        color = status_colors[status]
        # Background bar
        ax.barh(i, 100, height=0.65, color="#E8E8E8", zorder=1)
        # Progress bar
        if pct > 0:
            ax.barh(i, pct, height=0.65, color=color, zorder=2, edgecolor="white", linewidth=0.5)
        # Percentage label
        ax.text(pct + 1.5, i, f"{pct}%", va="center", fontsize=8, fontweight="bold",
                color=DARK_GRAY, zorder=3)
        # Status badge
        badge_color = status_colors[status]
        ax.text(103, i, f"  {status}  ", va="center", fontsize=7, fontweight="bold",
                color=status_text_colors[status],
                bbox=dict(boxstyle="round,pad=0.3", facecolor=badge_color, edgecolor="none", alpha=0.9))

    ax.set_yticks(y_pos)
    ax.set_yticklabels(names, fontsize=8, fontweight="bold")
    ax.set_xlim(0, 130)
    ax.set_xlabel("Completion %", fontsize=10, color=DARK_GRAY)
    ax.set_title("Phase 1 Deliverables Progress Tracker", fontsize=13, fontweight="bold",
                 color=DEEP_BLUE, pad=12)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(axis="x", alpha=0.15)
    fig.tight_layout()
    return save_figure(fig, "03_deliverables_progress")


# ---------------------------------------------------------------------------
# Diagram 4: Phase Dependency Flow
# ---------------------------------------------------------------------------

def create_dependency_flow():
    fig, ax = plt.subplots(figsize=(14, 5))
    fig.set_facecolor(WHITE)
    ax.set_facecolor(WHITE)

    phases = [
        ("Phase 1\nFoundation\n$115K", 1, 2.5, DEEP_BLUE),
        ("Phase 2\nIntegration\n$150K", 4, 2.5, TEAL),
        ("Phase 3\nOptimization\n$80K", 7, 2.5, SUCCESS_GREEN),
        ("Phase 4\nAutomation\n$100K", 10, 2.5, PURPLE),
        ("Phase 5\nInnovation\n$85K", 13, 2.5, ACCENT_GOLD),
    ]

    sub_items = [
        (1, 0.6, "AWS Setup\nTerraform\nFastAPI Core", DEEP_BLUE),
        (4, 0.6, "ServiceNow\nEthos API\nMigration", TEAL),
        (7, 0.6, "Scaling\nPerformance\nDR", SUCCESS_GREEN),
        (10, 0.6, "Self-Service\nAuto-Remediation\nIaC Library", PURPLE),
        (13, 0.6, "AI/ML\nMulti-Region\nSOC 2", ACCENT_GOLD),
    ]

    for label, x, y, color in phases:
        box = FancyBboxPatch((x - 1.1, y - 0.7), 2.2, 1.4,
                             boxstyle="round,pad=0.15", facecolor=color,
                             edgecolor="white", linewidth=2, zorder=3)
        ax.add_patch(box)
        ax.text(x, y, label, ha="center", va="center", fontsize=8,
                fontweight="bold", color="white", zorder=4)

    for x, y, label, color in sub_items:
        box = FancyBboxPatch((x - 1.05, y - 0.55), 2.1, 1.0,
                             boxstyle="round,pad=0.1", facecolor=color, alpha=0.15,
                             edgecolor=color, linewidth=1, zorder=2)
        ax.add_patch(box)
        ax.text(x, y, label, ha="center", va="center", fontsize=6.5,
                color=color, fontweight="bold", zorder=3)

    # Vertical arrows from phase box to sub-item
    for x in [1, 4, 7, 10, 13]:
        ax.annotate("", xy=(x, 1.15), xytext=(x, 1.8),
                    arrowprops=dict(arrowstyle="-|>", color="#7F8C8D", lw=1.5, mutation_scale=12))

    # Horizontal arrows between phase boxes
    for i in range(4):
        x_start = phases[i][1] + 1.1
        x_end = phases[i + 1][1] - 1.1
        ax.annotate("", xy=(x_end, 2.5), xytext=(x_start, 2.5),
                    arrowprops=dict(arrowstyle="-|>", color=DARK_GRAY, lw=2, mutation_scale=15))

    ax.set_xlim(-1, 15)
    ax.set_ylim(-0.3, 4.3)
    ax.set_title("Phase Dependency Flow Diagram", fontsize=14, fontweight="bold",
                 color=DEEP_BLUE, pad=12)
    ax.axis("off")
    fig.tight_layout()
    return save_figure(fig, "04_dependency_flow")


# ---------------------------------------------------------------------------
# Diagram 5: Year 2+ Innovation Radar
# ---------------------------------------------------------------------------

def create_innovation_radar():
    fig, ax = plt.subplots(figsize=(8, 8), subplot_kw=dict(polar=True))
    fig.set_facecolor(WHITE)

    categories = ["AI/ML", "IoT", "Blockchain", "Serverless", "Data Lake"]
    N = len(categories)

    biz_value = [9, 7, 6, 8, 8]
    tech_ready = [7, 5, 4, 8, 7]
    timeline = [8, 6, 5, 7, 6]

    angles = np.linspace(0, 2 * np.pi, N, endpoint=False).tolist()
    angles += angles[:1]

    for data, color, label in [
        (biz_value, DEEP_BLUE, "Business Value"),
        (tech_ready, TEAL, "Technical Readiness"),
        (timeline, ACCENT_GOLD, "Timeline Proximity"),
    ]:
        values = data + data[:1]
        ax.plot(angles, values, "o-", linewidth=2.5, color=color, label=label, markersize=7)
        ax.fill(angles, values, alpha=0.1, color=color)

    ax.set_xticks(angles[:-1])
    ax.set_xticklabels(categories, fontsize=11, fontweight="bold", color=DARK_GRAY)
    ax.set_ylim(0, 10)
    ax.set_yticks([2, 4, 6, 8, 10])
    ax.set_yticklabels(["2", "4", "6", "8", "10"], fontsize=8, color="#888")
    ax.set_rlabel_position(30)
    ax.spines["polar"].set_color("#DDD")
    ax.grid(color="#DDD", linewidth=0.5)

    ax.legend(loc="upper right", bbox_to_anchor=(1.35, 1.15), fontsize=10,
              frameon=True, fancybox=True, shadow=True)
    ax.set_title("Year 2+ Innovation Radar (2028+)", fontsize=14, fontweight="bold",
                 color=DEEP_BLUE, pad=24)
    fig.tight_layout()
    return save_figure(fig, "05_innovation_radar")


# ---------------------------------------------------------------------------
# Diagram 6: Technology Stack Architecture (layered)
# ---------------------------------------------------------------------------

def create_tech_stack():
    fig, ax = plt.subplots(figsize=(12, 8))
    fig.set_facecolor(WHITE)
    ax.set_facecolor(WHITE)

    layers = [
        ("Infrastructure Layer\nAWS VPC  |  ECS Fargate  |  ALB  |  Route 53  |  CloudWatch", 0, "#34495E"),
        ("Data Layer\nPostgreSQL RDS  |  Redis ElastiCache  |  S3  |  Secrets Manager", 1.5, DEEP_BLUE),
        ("Service Layer\nServiceNow ITSM  |  Ellucian Ethos  |  Email  |  Monitoring", 3.0, TEAL),
        ("Application Layer\nFastAPI  |  Pydantic  |  SQLAlchemy  |  Celery  |  JWT Auth", 4.5, SUCCESS_GREEN),
        ("API Gateway\nAWS API Gateway  |  Rate Limiting  |  WAF  |  SSL/TLS", 6.0, PURPLE),
        ("Frontend / Consumer Layer\nReact Dashboard  |  ServiceNow Portal  |  Mobile  |  CLI Tools", 7.5, ACCENT_GOLD),
    ]

    for label, y, color in layers:
        box = FancyBboxPatch((0.5, y), 11, 1.2, boxstyle="round,pad=0.15",
                             facecolor=color, edgecolor="white", linewidth=2,
                             alpha=0.9, zorder=3)
        ax.add_patch(box)
        ax.text(6, y + 0.6, label, ha="center", va="center",
                fontsize=10, fontweight="bold", color="white", zorder=4,
                linespacing=1.5)

    # Arrows between layers
    for i in range(len(layers) - 1):
        y_bottom = layers[i][1] + 1.2
        y_top = layers[i + 1][1]
        mid = (y_bottom + y_top) / 2
        ax.annotate("", xy=(6, y_top), xytext=(6, y_bottom),
                    arrowprops=dict(arrowstyle="-|>", color="#BDC3C7", lw=2, mutation_scale=18))

    # Side labels
    cross_labels = [
        ("Security\n& IAM", -0.1, 4.5, ALERT_RED),
        ("CI/CD\nPipeline", 12.1, 4.5, "#2C3E50"),
        ("Monitoring\n& Logging", -0.1, 1.5, "#2C3E50"),
        ("Terraform\nIaC", 12.1, 1.5, ALERT_RED),
    ]
    for label, x, y, color in cross_labels:
        ax.text(x, y, label, ha="center", va="center", fontsize=8,
                fontweight="bold", color=color, style="italic",
                bbox=dict(boxstyle="round,pad=0.4", facecolor="#F0F0F0",
                          edgecolor=color, linewidth=1))

    ax.set_xlim(-1.5, 13.5)
    ax.set_ylim(-0.5, 9.5)
    ax.set_title("ECTP Technology Stack Architecture", fontsize=14, fontweight="bold",
                 color=DEEP_BLUE, pad=12)
    ax.axis("off")
    fig.tight_layout()
    return save_figure(fig, "06_tech_stack")


# ---------------------------------------------------------------------------
# Diagram 7: Success Metrics Dashboard (multi-panel)
# ---------------------------------------------------------------------------

def create_metrics_dashboard():
    fig, axes = plt.subplots(2, 3, figsize=(14, 8))
    fig.set_facecolor(WHITE)
    fig.suptitle("Success Metrics Dashboard  --  Targets by Phase", fontsize=15,
                 fontweight="bold", color=DEEP_BLUE, y=1.0)

    phase_labels = ["P1", "P2", "P3", "P4", "P5"]
    phase_colors = [DEEP_BLUE, TEAL, SUCCESS_GREEN, PURPLE, ACCENT_GOLD]

    # 1) Uptime %
    ax = axes[0, 0]
    vals = [99.0, 99.5, 99.9, 99.9, 99.95]
    bars = ax.bar(phase_labels, vals, color=phase_colors, edgecolor="white", linewidth=1)
    ax.set_ylim(98.5, 100.1)
    ax.set_title("Uptime Target (%)", fontsize=10, fontweight="bold", color=DEEP_BLUE)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width()/2, v + 0.05, f"{v}%", ha="center", fontsize=8, fontweight="bold")
    ax.set_facecolor("#F8FAFC")

    # 2) Latency p99 (ms)
    ax = axes[0, 1]
    vals = [500, 500, 300, 300, 200]
    bars = ax.bar(phase_labels, vals, color=phase_colors, edgecolor="white", linewidth=1)
    ax.set_title("p99 Latency Target (ms)", fontsize=10, fontweight="bold", color=DEEP_BLUE)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width()/2, v + 10, f"{v}ms", ha="center", fontsize=8, fontweight="bold")
    ax.set_facecolor("#F8FAFC")

    # 3) Cost Reduction %
    ax = axes[0, 2]
    vals = [0, 5, 25, 30, 40]
    bars = ax.bar(phase_labels, vals, color=phase_colors, edgecolor="white", linewidth=1)
    ax.set_title("Cumulative Cost Reduction (%)", fontsize=10, fontweight="bold", color=DEEP_BLUE)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width()/2, v + 1, f"{v}%", ha="center", fontsize=8, fontweight="bold")
    ax.set_facecolor("#F8FAFC")

    # 4) CI/CD Build Time (min)
    ax = axes[1, 0]
    vals = [15, 12, 10, 8, 5]
    bars = ax.bar(phase_labels, vals, color=phase_colors, edgecolor="white", linewidth=1)
    ax.set_title("CI/CD Build Time (min)", fontsize=10, fontweight="bold", color=DEEP_BLUE)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width()/2, v + 0.3, f"{v}m", ha="center", fontsize=8, fontweight="bold")
    ax.set_facecolor("#F8FAFC")

    # 5) RTO (hours)
    ax = axes[1, 1]
    vals = [24, 12, 4, 2, 0.5]
    bars = ax.bar(phase_labels, vals, color=phase_colors, edgecolor="white", linewidth=1)
    ax.set_title("Recovery Time Objective (hrs)", fontsize=10, fontweight="bold", color=DEEP_BLUE)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width()/2, v + 0.5, f"{v}h", ha="center", fontsize=8, fontweight="bold")
    ax.set_facecolor("#F8FAFC")

    # 6) Automation %
    ax = axes[1, 2]
    vals = [10, 25, 40, 60, 80]
    bars = ax.bar(phase_labels, vals, color=phase_colors, edgecolor="white", linewidth=1)
    ax.set_title("Automation Coverage (%)", fontsize=10, fontweight="bold", color=DEEP_BLUE)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width()/2, v + 1.5, f"{v}%", ha="center", fontsize=8, fontweight="bold")
    ax.set_facecolor("#F8FAFC")

    for row in axes:
        for ax in row:
            ax.spines["top"].set_visible(False)
            ax.spines["right"].set_visible(False)
            ax.grid(axis="y", alpha=0.15)

    fig.tight_layout(rect=[0, 0, 1, 0.96])
    return save_figure(fig, "07_metrics_dashboard")


# =====================================================================
# Word Document Helpers
# =====================================================================

def set_cell_shading(cell, color_hex):
    """Set background color of a table cell."""
    color_hex = color_hex.lstrip("#")
    shading_elm = parse_xml(
        f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>'
    )
    cell._element.get_or_add_tcPr().append(shading_elm)


def set_cell_text(cell, text, bold=False, color=None, size=9, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    """Write styled text into a table cell."""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = alignment
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color
    # Reduce cell padding
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)


def add_styled_table(doc, headers, rows, col_widths=None):
    """Add a table with colored header and alternating row shading."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Header row
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        set_cell_shading(cell, DEEP_BLUE)
        set_cell_text(cell, h, bold=True, color=WHITE_RGB, size=9,
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Data rows
    for i, row_data in enumerate(rows):
        bg = LIGHT_BLUE if i % 2 == 0 else WHITE
        for j, value in enumerate(row_data):
            cell = table.rows[i + 1].cells[j]
            set_cell_shading(cell, bg)
            set_cell_text(cell, value, size=8)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Column widths
    if col_widths:
        for row in table.rows:
            for j, w in enumerate(col_widths):
                row.cells[j].width = Inches(w)

    return table


def add_heading_with_bg(doc, text, level=1, color_hex=DEEP_BLUE):
    """Add a heading and colour its background via shading."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = WHITE_RGB
        run.font.name = "Cambria"
    # Apply shading to paragraph
    pPr = heading._element.get_or_add_pPr()
    shading = parse_xml(
        f'<w:shd {nsdecls("w")} w:fill="{color_hex.lstrip("#")}" w:val="clear"/>'
    )
    pPr.append(shading)
    fmt = heading.paragraph_format
    fmt.space_before = Pt(12)
    fmt.space_after = Pt(6)
    return heading


def add_callout_box(doc, title, text, color_hex=TEAL):
    """Add a colored callout / insight box."""
    p = doc.add_paragraph()
    pPr = p._element.get_or_add_pPr()
    shading = parse_xml(
        f'<w:shd {nsdecls("w")} w:fill="{color_hex.lstrip("#")}22" w:val="clear"/>'
    )
    pPr.append(shading)
    # Border
    border_xml = (
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="8" w:color="{color_hex.lstrip("#")}"/>'
        f'</w:pBdr>'
    )
    pPr.append(parse_xml(border_xml))

    run_title = p.add_run(f"{title}: ")
    run_title.bold = True
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = RGBColor(
        int(color_hex.lstrip("#")[:2], 16),
        int(color_hex.lstrip("#")[2:4], 16),
        int(color_hex.lstrip("#")[4:6], 16),
    )
    run_title.font.name = "Calibri"
    run_body = p.add_run(text)
    run_body.font.size = Pt(9)
    run_body.font.name = "Calibri"
    run_body.font.color.rgb = DARK_GRAY_RGB
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)


def add_body(doc, text, bold=False, size=10):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = DARK_GRAY_RGB
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return p


def add_bullet(doc, text, level=0, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.size = Pt(9)
        r.font.name = "Calibri"
        r.font.color.rgb = DARK_GRAY_RGB
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.font.name = "Calibri"
    r.font.color.rgb = DARK_GRAY_RGB
    if level > 0:
        p.paragraph_format.left_indent = Inches(0.5 * level)
    p.paragraph_format.space_after = Pt(2)
    return p


def add_image(doc, path, width=Inches(6.2)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=width)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)


def add_page_break(doc):
    doc.add_page_break()


# =====================================================================
# MAIN DOCUMENT BUILDER
# =====================================================================

def build_document():
    print("Generating diagrams...")
    img_timeline = create_roadmap_timeline()
    img_budget = create_budget_donut()
    img_deliverables = create_deliverables_progress()
    img_dependency = create_dependency_flow()
    img_radar = create_innovation_radar()
    img_tech_stack = create_tech_stack()
    img_metrics = create_metrics_dashboard()
    print("  All 7 diagrams created.")

    doc = Document()

    # -- Default styles --
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.font.color.rgb = DARK_GRAY_RGB
    style.paragraph_format.space_after = Pt(4)

    for level in range(1, 4):
        hs = doc.styles[f"Heading {level}"]
        hs.font.name = "Cambria"
        hs.font.color.rgb = DEEP_BLUE_RGB

    # -- Page setup --
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # -- Headers / Footers --
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("ECTP-ARCH-002  |  Architecture & Future Roadmap  |  CONFIDENTIAL")
    hr.font.size = Pt(7)
    hr.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    hr.font.name = "Calibri"

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Page number field
    fr = fp.add_run("Page ")
    fr.font.size = Pt(7)
    fr.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    fr._element.append(fldChar1)
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fr._element.append(instrText)
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    fr._element.append(fldChar2)
    fr2 = fp.add_run("  |  Enterprise Cloud Transformation Platform")
    fr2.font.size = Pt(7)
    fr2.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    # ===================================================================
    # COVER PAGE
    # ===================================================================
    print("Building cover page...")
    # Spacer
    for _ in range(4):
        doc.add_paragraph()

    # Title block
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Enterprise Cloud Transformation Platform")
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = DEEP_BLUE_RGB
    r.font.name = "Cambria"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("(ECTP)")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = TEAL_RGB
    r.font.name = "Cambria"

    # Decorative line
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("_" * 60)
    r.font.color.rgb = GOLD_RGB
    r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Architecture & Future Roadmap")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = DEEP_BLUE_RGB
    r.font.name = "Cambria"

    doc.add_paragraph()

    # Meta table
    meta_data = [
        ("Document ID", "ECTP-ARCH-002"),
        ("Version", "2.0.0"),
        ("Author", "Gopi Krishna Vajrala"),
        ("Date", "February 16, 2026"),
        ("Classification", "Internal - Confidential"),
        ("Status", "Active"),
    ]
    meta_table = doc.add_table(rows=len(meta_data), cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (key, val) in enumerate(meta_data):
        set_cell_shading(meta_table.rows[i].cells[0], DEEP_BLUE)
        set_cell_text(meta_table.rows[i].cells[0], key, bold=True, color=WHITE_RGB, size=10,
                      alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        meta_table.rows[i].cells[0].width = Inches(2.0)
        set_cell_shading(meta_table.rows[i].cells[1], LIGHT_BLUE)
        set_cell_text(meta_table.rows[i].cells[1], val, bold=False, color=DEEP_BLUE_RGB, size=10)
        meta_table.rows[i].cells[1].width = Inches(3.5)

    add_page_break(doc)

    # ===================================================================
    # TABLE OF CONTENTS
    # ===================================================================
    print("Adding Table of Contents...")
    add_heading_with_bg(doc, "Table of Contents", level=1, color_hex=DEEP_BLUE)
    toc_items = [
        ("1", "Roadmap Overview"),
        ("2", "Phase 1 -- Foundation & Core Platform"),
        ("3", "Phase 2 -- Integration & Migration"),
        ("4", "Phase 3 -- Optimization & Scale"),
        ("5", "Phase 4 -- Advanced Automation"),
        ("6", "Phase 5 -- Innovation & Enhancement"),
        ("7", "Year 2+ Innovation Roadmap (2028+)"),
        ("8", "Architecture Decision Records (ADRs)"),
        ("9", "Roadmap Governance"),
    ]
    for num, title in toc_items:
        p = doc.add_paragraph()
        r1 = p.add_run(f"Section {num}:  ")
        r1.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = DEEP_BLUE_RGB
        r1.font.name = "Calibri"
        r2 = p.add_run(title)
        r2.font.size = Pt(11)
        r2.font.color.rgb = DARK_GRAY_RGB
        r2.font.name = "Calibri"
        p.paragraph_format.space_after = Pt(3)

    add_page_break(doc)

    # ===================================================================
    # SECTION 1: ROADMAP OVERVIEW
    # ===================================================================
    print("Section 1: Roadmap Overview...")
    add_heading_with_bg(doc, "1. Roadmap Overview", level=1, color_hex=DEEP_BLUE)

    add_body(doc, (
        "The Enterprise Cloud Transformation Platform (ECTP) roadmap defines a comprehensive "
        "5-phase journey spanning January 2026 through December 2027, with a total investment "
        "of $530,000. This roadmap transforms the university's IT operations from manual, "
        "on-premises processes to a fully automated, cloud-native platform built on AWS."
    ))

    add_callout_box(doc, "Key Insight",
                    "The $530K investment is projected to deliver a 3x ROI within 3 years through "
                    "operational efficiency gains, reduced downtime, and infrastructure cost optimization.",
                    TEAL)

    # Phase summary table
    add_heading_with_bg(doc, "Phase Summary", level=2, color_hex=TEAL)
    add_styled_table(doc,
        ["Phase", "Name", "Timeline", "Budget", "Key Focus"],
        [
            ["1", "Foundation & Core Platform", "Jan - Apr 2026", "$115,000", "AWS setup, FastAPI, Terraform, CI/CD"],
            ["2", "Integration & Migration", "May - Sep 2026", "$150,000", "ServiceNow, Ethos, workload migration"],
            ["3", "Optimization & Scale", "Oct 2026 - Jan 2027", "$80,000", "Performance, DR, compliance"],
            ["4", "Advanced Automation", "Feb - Jun 2027", "$100,000", "Self-service, auto-remediation, IaC library"],
            ["5", "Innovation & Enhancement", "Jul - Dec 2027", "$85,000", "AI/ML, multi-region, SOC 2 Type II"],
        ],
        col_widths=[0.5, 1.6, 1.3, 0.8, 2.5],
    )

    doc.add_paragraph()
    add_image(doc, img_timeline, width=Inches(6.5))

    doc.add_paragraph()
    add_image(doc, img_budget, width=Inches(5.5))

    add_page_break(doc)

    # ===================================================================
    # SECTION 2: PHASE 1 - FOUNDATION
    # ===================================================================
    print("Section 2: Phase 1 - Foundation...")
    add_heading_with_bg(doc, "2. Phase 1 -- Foundation & Core Platform", level=1, color_hex=DEEP_BLUE)

    # Phase metadata
    phase1_meta = doc.add_table(rows=1, cols=4)
    phase1_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (label, val, bg) in enumerate([
        ("Timeline", "Jan - Apr 2026", TEAL),
        ("Budget", "$115,000", SUCCESS_GREEN),
        ("Status", "In Progress", ACCENT_GOLD),
        ("Lead", "Gopi Krishna Vajrala", PURPLE),
    ]):
        cell = phase1_meta.rows[0].cells[j]
        set_cell_shading(cell, bg)
        set_cell_text(cell, f"{label}: {val}", bold=True, color=WHITE_RGB, size=9,
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph()
    add_body(doc, (
        "Phase 1 establishes the foundational infrastructure and core platform components. "
        "This phase focuses on setting up the AWS environment, building the core FastAPI application, "
        "implementing Infrastructure as Code with Terraform, and establishing CI/CD pipelines."
    ))

    # Deliverables table
    add_heading_with_bg(doc, "Phase 1 Deliverables", level=2, color_hex=TEAL)
    add_styled_table(doc,
        ["ID", "Deliverable", "Owner", "Target Date", "Status"],
        [
            ["1.1", "Architecture design and ADRs", "Gopi", "Jan 31, 2026", "Complete"],
            ["1.2", "Core FastAPI platform", "Dev Team", "Feb 15, 2026", "Complete"],
            ["1.3", "Health check endpoints", "Dev Team", "Feb 15, 2026", "Complete"],
            ["1.4", "Terraform modules (VPC, ECS, RDS)", "Cloud Architect", "Feb 28, 2026", "In Progress"],
            ["1.5", "CI/CD pipeline (GitHub Actions)", "DevOps", "Mar 15, 2026", "Planned"],
            ["1.6", "Docker multi-stage build & ECR", "DevOps", "Mar 15, 2026", "Planned"],
            ["1.7", "ServiceNow ITSM integration", "Integration Team", "Mar 31, 2026", "Planned"],
            ["1.8", "Ellucian Ethos API integration", "Integration Team", "Mar 31, 2026", "Planned"],
            ["1.9", "Security baseline & scanning", "Security Lead", "Apr 15, 2026", "Planned"],
            ["1.10", "Production deployment", "Full Team", "Apr 30, 2026", "Planned"],
        ],
        col_widths=[0.5, 2.2, 1.2, 1.1, 0.9],
    )

    doc.add_paragraph()
    add_image(doc, img_deliverables, width=Inches(6.3))

    # Dependencies
    add_heading_with_bg(doc, "Dependencies", level=2, color_hex=TEAL)
    deps = [
        ("AWS Account Provisioning: ", "Complete -- all four environments (dev, staging, UAT, prod) provisioned."),
        ("ServiceNow API Access: ", "In Progress -- credentials requested, awaiting approval from ITSM team."),
        ("Ellucian Ethos Credentials: ", "In Progress -- API key request submitted, expected by March 1."),
        ("Security Review Scheduling: ", "Pending -- initial assessment scheduled for March 15."),
    ]
    for prefix, text in deps:
        add_bullet(doc, text, bold_prefix=prefix)

    # Risks
    add_heading_with_bg(doc, "Risks & Mitigations", level=2, color_hex=ALERT_RED)
    add_styled_table(doc,
        ["Risk", "Impact", "Probability", "Mitigation"],
        [
            ["ServiceNow API delays", "High", "Medium",
             "Develop mock API layer; parallel development track"],
            ["Ellucian credential provisioning", "High", "Medium",
             "Escalation path defined; interim mock services"],
            ["Security review bottleneck", "Medium", "Low",
             "Pre-assessment checklist; early engagement with security team"],
        ],
        col_widths=[1.5, 0.7, 0.8, 3.0],
    )

    # Success Metrics
    doc.add_paragraph()
    add_heading_with_bg(doc, "Success Metrics", level=2, color_hex=SUCCESS_GREEN)
    p1_metrics = [
        "All 4 environments (dev, staging, UAT, prod) deployed and accessible",
        "Health check endpoints returning HTTP 200 across all environments",
        "CI/CD pipeline completing full build-test-deploy within 15 minutes",
        "Terraform plan validates with zero errors on all modules",
        "Zero critical or high security findings in baseline scan",
    ]
    for m in p1_metrics:
        add_bullet(doc, m)

    add_callout_box(doc, "Key Insight",
                    "Phase 1 is currently on track with 3 of 10 deliverables complete and Terraform "
                    "modules actively in development. The critical path runs through ServiceNow and "
                    "Ellucian integrations.",
                    SUCCESS_GREEN)

    add_page_break(doc)

    # ===================================================================
    # SECTION 3: PHASE 2 - INTEGRATION & MIGRATION
    # ===================================================================
    print("Section 3: Phase 2 - Integration & Migration...")
    add_heading_with_bg(doc, "3. Phase 2 -- Integration & Migration", level=1, color_hex=DEEP_BLUE)

    phase2_meta = doc.add_table(rows=1, cols=4)
    phase2_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (label, val, bg) in enumerate([
        ("Timeline", "May - Sep 2026", TEAL),
        ("Budget", "$150,000", SUCCESS_GREEN),
        ("Status", "Planned", "#7F8C8D"),
        ("Lead", "Integration Team", PURPLE),
    ]):
        cell = phase2_meta.rows[0].cells[j]
        set_cell_shading(cell, bg)
        set_cell_text(cell, f"{label}: {val}", bold=True, color=WHITE_RGB, size=9,
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph()
    add_body(doc, (
        "Phase 2 focuses on deep integration with enterprise systems and beginning the migration "
        "of existing workloads to the cloud platform. This is the largest budget phase reflecting "
        "the complexity of integrating ServiceNow, Ellucian Ethos, and migrating production workloads."
    ))

    add_heading_with_bg(doc, "Phase 2 Deliverables", level=2, color_hex=TEAL)
    add_styled_table(doc,
        ["ID", "Deliverable", "Owner", "Target Date", "Status"],
        [
            ["2.1", "ServiceNow full ITSM integration", "Integration Team", "May 31, 2026", "Planned"],
            ["2.2", "Ellucian Ethos data sync pipeline", "Integration Team", "Jun 15, 2026", "Planned"],
            ["2.3", "Workload migration - Batch 1 (2 apps)", "Cloud Architect", "Jun 30, 2026", "Planned"],
            ["2.4", "Workload migration - Batch 2 (3 apps)", "Cloud Architect", "Jul 31, 2026", "Planned"],
            ["2.5", "API gateway with rate limiting", "Dev Team", "Jun 30, 2026", "Planned"],
            ["2.6", "Centralized logging (CloudWatch/ELK)", "DevOps", "Jul 15, 2026", "Planned"],
            ["2.7", "Cost management dashboard", "Cloud Architect", "Aug 15, 2026", "Planned"],
            ["2.8", "Load testing framework", "QA Team", "Aug 31, 2026", "Planned"],
            ["2.9", "Data migration validation suite", "Data Team", "Sep 15, 2026", "Planned"],
            ["2.10", "Phase 2 security assessment", "Security Lead", "Sep 30, 2026", "Planned"],
        ],
        col_widths=[0.5, 2.2, 1.2, 1.1, 0.9],
    )

    doc.add_paragraph()
    add_heading_with_bg(doc, "Success Metrics", level=2, color_hex=SUCCESS_GREEN)
    p2_metrics = [
        "ServiceNow API response time < 2 seconds for all ITSM operations",
        "Ellucian Ethos data synchronization completing within 15 minutes",
        "2 workloads successfully migrated with zero data loss",
        "Cost management dashboard achieving 5% budget accuracy",
        "Load test sustaining 10x normal traffic with p99 latency < 500ms",
    ]
    for m in p2_metrics:
        add_bullet(doc, m)

    add_callout_box(doc, "Key Insight",
                    "Phase 2 represents the highest investment ($150K) due to the complexity "
                    "of enterprise integrations and the risk associated with production workload migrations. "
                    "Success here unlocks the value for all subsequent phases.",
                    TEAL)

    add_page_break(doc)

    # ===================================================================
    # SECTION 4: PHASE 3 - OPTIMIZATION & SCALE
    # ===================================================================
    print("Section 4: Phase 3 - Optimization & Scale...")
    add_heading_with_bg(doc, "4. Phase 3 -- Optimization & Scale", level=1, color_hex=DEEP_BLUE)

    phase3_meta = doc.add_table(rows=1, cols=4)
    phase3_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (label, val, bg) in enumerate([
        ("Timeline", "Oct 2026 - Jan 2027", TEAL),
        ("Budget", "$80,000", SUCCESS_GREEN),
        ("Status", "Planned", "#7F8C8D"),
        ("Lead", "Cloud Architect", PURPLE),
    ]):
        cell = phase3_meta.rows[0].cells[j]
        set_cell_shading(cell, bg)
        set_cell_text(cell, f"{label}: {val}", bold=True, color=WHITE_RGB, size=9,
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph()
    add_body(doc, (
        "Phase 3 focuses on optimizing the platform for peak performance during critical periods "
        "like enrollment, establishing disaster recovery capabilities, and beginning the compliance "
        "journey toward SOC 2 Type I certification."
    ))

    add_heading_with_bg(doc, "Phase 3 Deliverables", level=2, color_hex=TEAL)
    add_styled_table(doc,
        ["ID", "Deliverable", "Owner", "Target Date", "Status"],
        [
            ["3.1", "Auto-scaling policies (ECS, RDS)", "Cloud Architect", "Oct 31, 2026", "Planned"],
            ["3.2", "Performance optimization & caching", "Dev Team", "Nov 15, 2026", "Planned"],
            ["3.3", "Disaster recovery (multi-AZ)", "Cloud Architect", "Nov 30, 2026", "Planned"],
            ["3.4", "Cost optimization (Reserved/Spot)", "Cloud Architect", "Dec 15, 2026", "Planned"],
            ["3.5", "Advanced monitoring & alerting", "DevOps", "Dec 31, 2026", "Planned"],
            ["3.6", "SOC 2 Type I preparation", "Security Lead", "Jan 15, 2027", "Planned"],
            ["3.7", "Enrollment peak load optimization", "Full Team", "Jan 15, 2027", "Planned"],
            ["3.8", "Phase 3 compliance audit", "Security Lead", "Jan 31, 2027", "Planned"],
        ],
        col_widths=[0.5, 2.2, 1.2, 1.1, 0.9],
    )

    doc.add_paragraph()
    add_heading_with_bg(doc, "Success Metrics", level=2, color_hex=SUCCESS_GREEN)
    p3_metrics = [
        "99.9% uptime maintained during enrollment peak periods",
        "p99 latency < 300ms under peak load conditions",
        "25% reduction in infrastructure costs through optimization",
        "4-hour Recovery Time Objective (RTO) for disaster recovery",
        "SOC 2 Type I readiness confirmed by auditor pre-assessment",
    ]
    for m in p3_metrics:
        add_bullet(doc, m)

    add_page_break(doc)

    # ===================================================================
    # SECTION 5: PHASE 4 - ADVANCED AUTOMATION
    # ===================================================================
    print("Section 5: Phase 4 - Advanced Automation...")
    add_heading_with_bg(doc, "5. Phase 4 -- Advanced Automation", level=1, color_hex=DEEP_BLUE)

    phase4_meta = doc.add_table(rows=1, cols=4)
    phase4_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (label, val, bg) in enumerate([
        ("Timeline", "Feb - Jun 2027", TEAL),
        ("Budget", "$100,000", SUCCESS_GREEN),
        ("Status", "Planned", "#7F8C8D"),
        ("Lead", "DevOps Lead", PURPLE),
    ]):
        cell = phase4_meta.rows[0].cells[j]
        set_cell_shading(cell, bg)
        set_cell_text(cell, f"{label}: {val}", bold=True, color=WHITE_RGB, size=9,
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph()
    add_body(doc, (
        "Phase 4 elevates the platform through advanced automation capabilities including "
        "self-service portals, automated incident remediation, infrastructure drift detection, "
        "and a shared Terraform module library for cross-team adoption."
    ))

    add_heading_with_bg(doc, "Phase 4 Deliverables", level=2, color_hex=TEAL)
    add_styled_table(doc,
        ["ID", "Deliverable", "Owner", "Target Date", "Status"],
        [
            ["4.1", "Self-service portal for common requests", "Dev Team", "Feb 28, 2027", "Planned"],
            ["4.2", "Automated incident remediation (runbooks)", "DevOps", "Mar 31, 2027", "Planned"],
            ["4.3", "Infrastructure drift detection & correction", "Cloud Architect", "Apr 15, 2027", "Planned"],
            ["4.4", "Shared Terraform module library", "Cloud Architect", "Apr 30, 2027", "Planned"],
            ["4.5", "ChatOps integration (Slack/Teams)", "Dev Team", "May 15, 2027", "Planned"],
            ["4.6", "Environment provisioning automation", "DevOps", "May 31, 2027", "Planned"],
            ["4.7", "Compliance-as-Code framework", "Security Lead", "Jun 15, 2027", "Planned"],
            ["4.8", "Phase 4 maturity assessment", "Full Team", "Jun 30, 2027", "Planned"],
        ],
        col_widths=[0.5, 2.3, 1.2, 1.0, 0.9],
    )

    doc.add_paragraph()
    add_heading_with_bg(doc, "Success Metrics", level=2, color_hex=SUCCESS_GREEN)
    p4_metrics = [
        "50% of routine requests fulfilled through self-service portal",
        "30% of incidents auto-remediated without human intervention",
        "Infrastructure drift detected within 1 hour of occurrence",
        "3+ teams actively using shared Terraform module library",
        "New environment provisioning completed in under 1 hour",
    ]
    for m in p4_metrics:
        add_bullet(doc, m)

    add_callout_box(doc, "Key Insight",
                    "Phase 4 automation capabilities are projected to reduce operational toil by 40%, "
                    "freeing engineering time for innovation work in Phase 5.",
                    PURPLE)

    add_page_break(doc)

    # ===================================================================
    # SECTION 6: PHASE 5 - INNOVATION & ENHANCEMENT
    # ===================================================================
    print("Section 6: Phase 5 - Innovation & Enhancement...")
    add_heading_with_bg(doc, "6. Phase 5 -- Innovation & Enhancement", level=1, color_hex=DEEP_BLUE)

    phase5_meta = doc.add_table(rows=1, cols=4)
    phase5_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (label, val, bg) in enumerate([
        ("Timeline", "Jul - Dec 2027", TEAL),
        ("Budget", "$85,000", SUCCESS_GREEN),
        ("Status", "Planned", "#7F8C8D"),
        ("Lead", "Full Team", PURPLE),
    ]):
        cell = phase5_meta.rows[0].cells[j]
        set_cell_shading(cell, bg)
        set_cell_text(cell, f"{label}: {val}", bold=True, color=WHITE_RGB, size=9,
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph()
    add_body(doc, (
        "Phase 5 represents the innovation horizon, introducing AI/ML capabilities, "
        "multi-region architecture, and achieving SOC 2 Type II certification. This phase "
        "transforms ECTP from an operational platform into a strategic innovation enabler."
    ))

    add_heading_with_bg(doc, "Phase 5 Deliverables", level=2, color_hex=TEAL)
    add_styled_table(doc,
        ["ID", "Deliverable", "Owner", "Target Date", "Status"],
        [
            ["5.1", "AI/ML-powered predictive scaling", "Dev Team", "Aug 31, 2027", "Planned"],
            ["5.2", "Multi-region active-passive failover", "Cloud Architect", "Sep 30, 2027", "Planned"],
            ["5.3", "SOC 2 Type II certification", "Security Lead", "Oct 31, 2027", "Planned"],
            ["5.4", "Event-driven architecture (10K events/min)", "Dev Team", "Nov 15, 2027", "Planned"],
            ["5.5", "Advanced analytics & reporting dashboard", "Data Team", "Nov 30, 2027", "Planned"],
            ["5.6", "Platform maturity Level 4 assessment", "Full Team", "Dec 31, 2027", "Planned"],
        ],
        col_widths=[0.5, 2.5, 1.2, 1.0, 0.9],
    )

    doc.add_paragraph()
    add_heading_with_bg(doc, "Success Metrics", level=2, color_hex=SUCCESS_GREEN)
    p5_metrics = [
        "Less than 2% change failure rate across all deployments",
        "Multi-region failover completing within 30 minutes",
        "SOC 2 Type II certification achieved",
        "Event-driven architecture processing 10,000 events per minute",
        "Platform maturity assessed at Level 4 (Managed/Optimizing)",
    ]
    for m in p5_metrics:
        add_bullet(doc, m)

    add_page_break(doc)

    # ===================================================================
    # DEPENDENCY FLOW & TECH STACK DIAGRAMS
    # ===================================================================
    print("Adding dependency flow and tech stack diagrams...")
    add_heading_with_bg(doc, "Phase Dependency Flow", level=2, color_hex=DEEP_BLUE)
    add_body(doc, (
        "The following diagram illustrates the dependencies between phases. Each phase builds "
        "upon the deliverables and capabilities established in the preceding phase. Critical "
        "path dependencies are highlighted to support resource planning and risk management."
    ))
    add_image(doc, img_dependency, width=Inches(6.5))

    doc.add_paragraph()
    add_heading_with_bg(doc, "Technology Stack Architecture", level=2, color_hex=DEEP_BLUE)
    add_body(doc, (
        "The ECTP technology stack follows a layered architecture pattern, providing clear "
        "separation of concerns, independent scalability, and well-defined integration points "
        "between components."
    ))
    add_image(doc, img_tech_stack, width=Inches(6.3))

    add_page_break(doc)

    # ===================================================================
    # SECTION 7: YEAR 2+ INNOVATION ROADMAP
    # ===================================================================
    print("Section 7: Year 2+ Innovation Roadmap...")
    add_heading_with_bg(doc, "7. Year 2+ Innovation Roadmap (2028+)", level=1, color_hex=DEEP_BLUE)

    add_body(doc, (
        "Beyond the initial 2-year roadmap, ECTP will evolve into a platform for institutional "
        "innovation. The following areas represent strategic investment opportunities identified "
        "through stakeholder interviews, industry analysis, and technology trend assessment."
    ))

    add_image(doc, img_radar, width=Inches(5.5))

    # AI/ML
    add_heading_with_bg(doc, "AI/ML Integration", level=2, color_hex=PURPLE)
    add_body(doc, "Five strategic AI/ML initiatives planned for 2028:")
    ai_items = [
        ("Predictive Auto-Scaling (Q1 2028): ", "ML models predict enrollment surges and scale infrastructure proactively."),
        ("Anomaly Detection (Q2 2028): ", "Real-time detection of performance anomalies and security threats."),
        ("Intelligent Incident Routing (Q2 2028): ", "NLP-based ticket classification and automatic routing to appropriate teams."),
        ("ChatBot Ops Assistant (Q3 2028): ", "Conversational AI for infrastructure queries and common operations."),
        ("Capacity Planning AI (Q4 2028): ", "Machine learning models for long-term capacity and cost forecasting."),
    ]
    for prefix, text in ai_items:
        add_bullet(doc, text, bold_prefix=prefix)

    # IoT
    add_heading_with_bg(doc, "IoT Integration", level=2, color_hex=TEAL)
    add_body(doc, "Four IoT initiatives for smart campus operations:")
    iot_items = [
        ("Smart Campus Monitoring (Q1 2028): ", "Environmental sensors for HVAC, lighting, and occupancy optimization."),
        ("Digital Signage Network (Q2 2028): ", "Centralized management of campus-wide digital displays."),
        ("Lab Equipment Monitoring (Q3 2028): ", "IoT sensors for lab equipment status, maintenance, and utilization."),
        ("Asset Tracking (Q4 2028): ", "RFID/BLE-based tracking for IT assets and equipment across campus."),
    ]
    for prefix, text in iot_items:
        add_bullet(doc, text, bold_prefix=prefix)

    # Blockchain
    add_heading_with_bg(doc, "Blockchain Credentials", level=2, color_hex=ACCENT_GOLD)
    add_body(doc, "Four blockchain-based credential initiatives:")
    bc_items = [
        ("Digital Diploma Verification (Q2 2028): ", "Blockchain-backed verifiable digital diplomas."),
        ("Micro-Credentials (Q3 2028): ", "Stackable, verifiable micro-credentials and digital badges."),
        ("Transcript Portability (Q4 2028): ", "Cross-institution transcript sharing via distributed ledger."),
        ("Alumni Verification (Q1 2029): ", "Instant alumni credential verification for employers."),
    ]
    for prefix, text in bc_items:
        add_bullet(doc, text, bold_prefix=prefix)

    # Additional areas
    add_heading_with_bg(doc, "Additional Strategic Initiatives", level=2, color_hex=DEEP_BLUE)
    additional = [
        ("Serverless Evolution: ", "Migration of suitable workloads to AWS Lambda for cost optimization and infinite scalability."),
        ("Data Lake / Lakehouse: ", "Centralized institutional data platform combining structured and unstructured data for advanced analytics."),
        ("Multi-Cloud Strategy: ", "Strategic evaluation of Azure and GCP for specific workloads to avoid vendor lock-in."),
        ("Edge Computing: ", "Campus edge nodes for latency-sensitive applications and local processing."),
        ("Zero Trust Architecture: ", "Full implementation of zero-trust security model with microsegmentation and continuous verification."),
    ]
    for prefix, text in additional:
        add_bullet(doc, text, bold_prefix=prefix)

    add_callout_box(doc, "Key Insight",
                    "Year 2+ initiatives are evaluated using the Innovation Radar framework, "
                    "scoring each opportunity on Business Value, Technical Readiness, and "
                    "Timeline Proximity. This ensures data-driven prioritization of innovation investments.",
                    PURPLE)

    add_page_break(doc)

    # ===================================================================
    # SECTION 8: ADRs
    # ===================================================================
    print("Section 8: Architecture Decision Records...")
    add_heading_with_bg(doc, "8. Architecture Decision Records (ADRs)", level=1, color_hex=DEEP_BLUE)

    add_body(doc, (
        "Architecture Decision Records capture the key technical decisions made during the ECTP "
        "project, including context, alternatives considered, and rationale for each decision. "
        "ADRs provide an auditable trail of architectural choices."
    ))

    add_heading_with_bg(doc, "ADR Register", level=2, color_hex=TEAL)
    add_styled_table(doc,
        ["ADR ID", "Title", "Status", "Date", "Decision"],
        [
            ["ADR-001", "Web Framework Selection", "Accepted", "Jan 15, 2026",
             "FastAPI for async performance and auto-generated OpenAPI docs"],
            ["ADR-002", "Primary Database", "Accepted", "Jan 15, 2026",
             "PostgreSQL on RDS for reliability, JSONB support, and managed service benefits"],
            ["ADR-003", "Container Orchestration", "Accepted", "Jan 18, 2026",
             "ECS Fargate for serverless containers eliminating node management overhead"],
            ["ADR-004", "Infrastructure as Code", "Accepted", "Jan 20, 2026",
             "Terraform for multi-provider support and mature module ecosystem"],
            ["ADR-005", "Caching Layer", "Accepted", "Jan 22, 2026",
             "Redis on ElastiCache for sub-millisecond latency and pub/sub capabilities"],
            ["ADR-006", "Data Validation", "Accepted", "Jan 25, 2026",
             "Pydantic for runtime validation, serialization, and settings management"],
        ],
        col_widths=[0.7, 1.3, 0.7, 0.9, 3.0],
    )

    doc.add_paragraph()
    add_heading_with_bg(doc, "ADR Process", level=2, color_hex=TEAL)
    add_body(doc, "The ADR lifecycle follows a structured review and approval process:")

    adr_process = [
        ("1. Draft: ", "Author creates ADR with context, alternatives, and recommendation."),
        ("2. Review (5 business days): ", "Team members review and provide feedback via pull request."),
        ("3. Discussion: ", "Architecture forum discusses trade-offs and concerns."),
        ("4. Decision (majority vote): ", "CCoE members vote on the recommendation."),
        ("5. Record: ", "Final decision is documented with rationale and consequences."),
        ("6. Communication: ", "Decision is communicated to all stakeholders via email and Confluence."),
    ]
    for prefix, text in adr_process:
        add_bullet(doc, text, bold_prefix=prefix)

    add_page_break(doc)

    # ===================================================================
    # SECTION 9: ROADMAP GOVERNANCE
    # ===================================================================
    print("Section 9: Roadmap Governance...")
    add_heading_with_bg(doc, "9. Roadmap Governance", level=1, color_hex=DEEP_BLUE)

    add_body(doc, (
        "Effective governance ensures the roadmap remains aligned with institutional strategy, "
        "adapts to changing requirements, and delivers measurable value. The governance framework "
        "covers review cadence, prioritization criteria, and change management."
    ))

    # Review Cadence
    add_heading_with_bg(doc, "Review Cadence", level=2, color_hex=TEAL)
    add_styled_table(doc,
        ["Review Type", "Frequency", "Participants", "Purpose"],
        [
            ["Sprint Review", "Bi-weekly", "Scrum Team, Product Owner",
             "Demonstrate completed work, gather feedback"],
            ["Phase Milestone", "At phase completion", "CCoE, Stakeholders",
             "Validate phase deliverables, approve next phase"],
            ["Roadmap Adjustment", "Quarterly", "CCoE, Governance Board",
             "Review priorities, adjust timeline/scope"],
            ["Innovation Review", "Semi-annually", "CCoE, Technology Leaders",
             "Evaluate emerging technologies, update Year 2+ roadmap"],
            ["Strategic Planning", "Annually", "Steering Committee, CIO",
             "Align roadmap with institutional strategy"],
        ],
        col_widths=[1.2, 1.1, 1.6, 2.7],
    )

    # Prioritization
    doc.add_paragraph()
    add_heading_with_bg(doc, "Prioritization Framework", level=2, color_hex=TEAL)
    add_body(doc, "Roadmap items are prioritized using a weighted scoring model across six dimensions:")
    prio_items = [
        ("Business Value (30%): ", "Direct impact on institutional operations, student experience, and revenue."),
        ("Technical Feasibility (20%): ", "Availability of skills, tools, and infrastructure to deliver."),
        ("Cost-Benefit (20%): ", "Expected ROI and payback period for the investment."),
        ("Risk (15%): ", "Technical, operational, and compliance risk assessment."),
        ("Dependencies (10%): ", "Blocking or enabling relationships with other initiatives."),
        ("Strategic Alignment (5%): ", "Fit with university's 5-year technology strategy."),
    ]
    for prefix, text in prio_items:
        add_bullet(doc, text, bold_prefix=prefix)

    # Change Process
    doc.add_paragraph()
    add_heading_with_bg(doc, "Change Management Process", level=2, color_hex=TEAL)
    add_body(doc, "Changes to the roadmap follow a structured evaluation and approval process:")

    change_steps = [
        ("1. ADR Proposal: ", "Requestor submits an ADR or change request with business justification."),
        ("2. CCoE Evaluation (2 weeks): ", "Cloud Center of Excellence evaluates technical and business impact."),
        ("3. Governance Board Review: ", "Board reviews CCoE recommendation and votes on inclusion."),
        ("4. Steering Committee (if >$50K): ", "Changes exceeding $50K require Steering Committee approval."),
        ("5. Add to Roadmap: ", "Approved changes are integrated into the roadmap with timeline and resources."),
        ("6. Communicate: ", "All stakeholders are notified of roadmap changes via established channels."),
    ]
    for prefix, text in change_steps:
        add_bullet(doc, text, bold_prefix=prefix)

    add_callout_box(doc, "Key Insight",
                    "The governance framework balances agility with oversight. Sprint-level reviews "
                    "enable rapid delivery while quarterly adjustments and annual planning ensure "
                    "strategic alignment. Changes under $50K can be approved within 2 weeks.",
                    DEEP_BLUE)

    # ===================================================================
    # METRICS DASHBOARD
    # ===================================================================
    add_page_break(doc)
    add_heading_with_bg(doc, "Success Metrics Dashboard -- All Phases", level=2, color_hex=DEEP_BLUE)
    add_body(doc, (
        "The following dashboard summarizes key performance targets across all five phases, "
        "illustrating the progressive improvement trajectory of the ECTP platform."
    ))
    add_image(doc, img_metrics, width=Inches(6.5))

    # ===================================================================
    # SAVE
    # ===================================================================
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    doc.save(OUTPUT_FILE)
    print(f"\nDocument saved to: {OUTPUT_FILE}")
    print(f"File size: {os.path.getsize(OUTPUT_FILE) / 1024:.1f} KB")

    # Cleanup temp images
    import shutil
    shutil.rmtree(TEMP_DIR, ignore_errors=True)
    print("Temporary diagram files cleaned up.")


if __name__ == "__main__":
    build_document()
